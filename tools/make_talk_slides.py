#!/usr/bin/env python3
"""문제점 → 개선 → 앞으로 예정.

문제점은 외부 발표 문장이 아니라, 현재 모델 계약
(고정 z_compact 32x64x64, 출력 262144, 기하/속성 분리, xyz detach)
에서 읽는다.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree
from pathlib import Path

OUT = "/data/daeho/aacd_proj/can3tok_encoder_decoder_new_fix_6/ED_문제점_개선_예정.pptx"
IMG = Path("/data/daeho/aacd_proj/can3tok_encoder_decoder_new_fix_6/assets/slides")

NAVY = RGBColor(0x1F, 0x2F, 0x4A)
RED = RGBColor(0xB8, 0x3B, 0x2E)
GREEN = RGBColor(0x1E, 0x6B, 0x48)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
LINE = RGBColor(0xDD, 0xDD, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xF4, 0xF6, 0xF8)
RED_BG = RGBColor(0xFD, 0xF0, 0xEE)
GREEN_BG = RGBColor(0xEA, 0xF5, 0xEF)

W, H = Inches(13.333), Inches(7.5)
FONT = "Malgun Gothic"
N = 8


def ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def fill(shp, c):
    shp.fill.solid()
    shp.fill.fore_color.rgb = c
    shp.line.fill.background()


def rect(s, x, y, w, h, c, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    fill(shp, c)
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    return shp


def txt(s, x, y, w, h, text, size=18, bold=False, color=INK, align="left", anchor="top"):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                 "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                   "right": PP_ALIGN.RIGHT}[align]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    ea(r)
    return box


def bullets(s, x, y, w, h, items, size=16, color=INK, after=8):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = "•  " + item
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
        ea(r)
    return box


def captioned_pic(s, path, x, y, w, h, cap):
    s.shapes.add_picture(str(path), x, y, w, h)
    rect(s, x, y + h, w, Inches(0.32), NAVY)
    txt(s, x, y + h, w, Inches(0.32), cap, 12, True, WHITE, "center", "middle")


def header(s, num_title, subtitle):
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, Inches(1.05), NAVY)
    txt(s, Inches(0.45), Inches(0.12), Inches(12.4), Inches(0.5),
        num_title, 26, True, WHITE, "left", "middle")
    txt(s, Inches(0.45), Inches(0.58), Inches(12.4), Inches(0.38),
        subtitle, 14, False, RGBColor(0xC8, 0xD2, 0xE0), "left", "middle")


def foot(s, i):
    txt(s, Inches(0.45), Inches(7.18), Inches(4), Inches(0.22),
        f"{i}  /  {N}", 11, False, MUTED)
    txt(s, Inches(8.5), Inches(7.18), Inches(4.4), Inches(0.22),
        "인코더–디코더  ·  문제점 → 개선 → 예정", 11, False, MUTED, "right")


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    GOLD = RGBColor(0xC9, 0x8A, 0x2C)

    # ---- 1 표지 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    txt(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.4),
        "can3tok  ·  고정 latent 32×64×64  ·  출력 262,144", 15, False,
        RGBColor(0xF0, 0xA8, 0x98))
    txt(s, Inches(0.7), Inches(2.15), Inches(12), Inches(1.15),
        "문제점  →  개선  →  앞으로 진행할 예정", 32, True, WHITE)
    txt(s, Inches(0.7), Inches(3.5), Inches(11.8), Inches(1.1),
        "구조상 디코더는 26.2만 개만 그리고, 렌더는 위치를 거의 못 움직인다.\n"
        "빈 표면을 큰 splat으로 메우는 길이 열리고, 그림이 흐려진다.",
        17, False, RGBColor(0xC5, 0xD0, 0xE0))
    for i, (lab, c) in enumerate([("1. 문제점", RED), ("2. 개선", GREEN), ("3. 예정", GOLD)]):
        x = Inches(0.7 + i * 3.3)
        rect(s, x, Inches(5.25), Inches(3.05), Inches(0.7), c)
        txt(s, x, Inches(5.25), Inches(3.05), Inches(0.7), lab, 18, True, WHITE, "center", "middle")
    foot(s, 1)

    # ---- 2 현재 구조 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "현재 모델이 고정한 계약", "월드모델이 먹는 것은 z_compact.  디코더 개수와 latent 크기는 안 늘린다.")
    foot(s, 2)
    stages = [
        ("입력", "원본 가우시안\nN  (가변)"),
        ("칸", "4096 앵커\n최근접 + spill 8"),
        ("인코더", "칸당 ≤512점\npool → 셀 코드"),
        ("latent", "z_compact\n32 × 64 × 64"),
        ("기하", "칸당 64 위치\nfolding 템플릿"),
        ("속성", "scale·rot\nopa·color"),
    ]
    bw, gap = Inches(1.95), Inches(0.18)
    x0 = Inches(0.4)
    for i, (t, b) in enumerate(stages):
        x = x0 + i * (bw + gap)
        rect(s, x, Inches(1.28), bw, Inches(1.85), PALE, LINE)
        rect(s, x, Inches(1.28), bw, Inches(0.38), NAVY)
        txt(s, x, Inches(1.28), bw, Inches(0.38), t, 13, True, WHITE, "center", "middle")
        txt(s, x + Inches(0.06), Inches(1.72), bw - Inches(0.12), Inches(1.3),
            b, 12, False, INK, "center", "middle")
        if i < len(stages) - 1:
            txt(s, x + bw - Inches(0.02), Inches(1.9), gap + Inches(0.08), Inches(0.4),
                "→", 18, True, MUTED, "center", "middle")
    txt(s, Inches(0.4), Inches(3.2), Inches(12.5), Inches(0.32),
        "출력은 항상 4096 × 64 = 262,144 splat.  인코더가 점을 더 봐도 디코더 개수는 그대로.",
        13, False, MUTED)

    cards = [
        (NAVY, "개수",
         "디코더 슬롯은 칸당 64개로 고정.\n원본이 77만·140만·200만이어도 그리는 점은 26.2만.\n남는 표면은 구조적으로 구멍이다."),
        (NAVY, "채널",
         "compact 32채널 = centroid 4 + occupancy 1\n+ shape 19 + appearance 8.\n칸 하나(64 splat)의 외형을 8채널에서 펼친다."),
        (NAVY, "그래디언트",
         "속성 디코더는 기하 xyz를 detach.\n렌더가 위치를 옮길 수 있는 양은\nnudge ≤ 0.15 × 그룹 크기 뿐."),
    ]
    for i, (c, t, b) in enumerate(cards):
        x = Inches(0.4 + i * 4.25)
        rect(s, x, Inches(3.62), Inches(4.05), Inches(3.15), PALE)
        rect(s, x, Inches(3.62), Inches(4.05), Inches(0.48), c)
        txt(s, x, Inches(3.62), Inches(4.05), Inches(0.48), t, 16, True, WHITE, "center", "middle")
        txt(s, x + Inches(0.18), Inches(4.22), Inches(3.7), Inches(2.4), b, 14, False, INK)

    # ---- 3 문제점 (구조에서) ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "1. 문제점", "위 계약이 학습에 열어 두는 세 길.  측정은 R8H (같은 디코더·같은 K).")
    foot(s, 3)
    probs = [
        (RED, "①  26.2만이 부분집합",
         "인코더 입력 cap을 올려도 디코더는 262,144만 방출한다.\n"
         "GT를 고른 부분집합으로 감독하면, 원본 N이 커질수록\n"
         "그 262k vs 사진 PSNR(canon 천장)이 내려간다.\n"
         "빈 칸은 나중에 splat이 덮을 자리이다."),
        (RED, "②  외형 8채널 / 64 splat",
         "appearance 코드 하나가 cross-attn으로 64 슬롯에 분배된다.\n"
         "위치 최근접으로 GT 속성을 붙이면, 같은 자리의\n"
         "가로·세로 얇은 가우시안이 평균되어 원이 된다.\n"
         "R8H에서 방향 다양성 = GT의 2.3%."),
        (RED, "③  렌더가 위치를 못 고침",
         "렌더 그래디언트는 위치와 cosine ≈ 0.13,\n"
         "점의 약 24%만 신호를 받는다 (측정).\n"
         "구멍은 점으로 메워지지 않고,\n"
         "이미 있는 점의 scale·opacity가 픽셀을 채운다."),
    ]
    for i, (c, t, b) in enumerate(probs):
        x = Inches(0.4 + i * 4.25)
        rect(s, x, Inches(1.28), Inches(4.05), Inches(5.5), RED_BG)
        rect(s, x, Inches(1.28), Inches(4.05), Inches(0.7), c)
        txt(s, x + Inches(0.1), Inches(1.28), Inches(3.85), Inches(0.7),
            t, 16, True, WHITE, "center", "middle")
        txt(s, x + Inches(0.2), Inches(2.15), Inches(3.65), Inches(4.4), b, 14, False, INK)

    # ---- 4 문제점 그림 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "1. 문제점  —  그림으로 보면", "같은 뷰.  왼쪽은 스냅샷 GT, 오른쪽은 디코더가 그린 자체 속성.")
    foot(s, 4)
    captioned_pic(s, IMG / "r8h_028630_orig.png", Inches(0.45), Inches(1.25),
                  Inches(6.05), Inches(3.38), "GT  ·  스냅샷 가우시안 렌더")
    captioned_pic(s, IMG / "r8h_028630_own.png", Inches(6.8), Inches(1.25),
                  Inches(6.05), Inches(3.38), "R8H  ·  예측 xyz + 예측 속성  (codec_own)")
    rect(s, Inches(0.45), Inches(5.05), Inches(12.4), Inches(1.85), PALE)
    bullets(s, Inches(0.65), Inches(5.18), Inches(12.0), Inches(1.6), [
        "글자·난간·바퀴가 큰 원형 splat에 먹힌다.  실루엣은 남고 고주파는 사라진다.",
        "R8H @10k 이 뷰: codec_own 19.03 dB / SSIM 0.59.  같은 xyz+GT 속성(codec)은 SSIM 0.67 — PSNR은 비슷해도 구조는 속성이 깎는다.",
        "전 평가: 사진 PSNR ≈ 16.5 dB, nn_unique ≈ 0.51.  예측의 절반이 같은 GT에 중복되고, 나머지 표면은 비어 있다.",
    ], 14, INK, 4)

    # ---- 5 원인 (구조 → blur) ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "1. 문제점  —  왜 큰 splat이 최적인가", "구조가 만든 구멍 + 목적함수가 구멍을 못 벌함.")
    foot(s, 5)
    bullets(s, Inches(0.5), Inches(1.22), Inches(12.3), Inches(1.85), [
        "Chamfer / Sinkhorn은 집합 거리라, 예측이 같은 GT에 여러 개 붙어도 손실이 작다.  빈 표면은 이 항에 거의 안 보인다.",
        "렌더는 ‘이 픽셀을 채워라’만 본다.  점을 빈 곳으로 옮기는 경로는 detach로 막혀 있고, log_scale 하나를 올리면 화면 타원이 바로 커진다.",
        "그래서 학습이 고르는 해는 ‘적은 점으로 넓게 칠하기’이다.  blur는 버그가 아니라 이 목적의 극소이다.",
    ], 15, INK, 6)
    rect(s, Inches(0.45), Inches(3.3), Inches(6.15), Inches(3.45), RED_BG)
    txt(s, Inches(0.65), Inches(3.45), Inches(5.8), Inches(0.38), "모델이 실제로 한 일", 16, True, RED)
    bullets(s, Inches(0.65), Inches(3.95), Inches(5.75), Inches(2.6), [
        "칸당 인코더 144점 → 밀집 영역 overflow 14–23%.  입력에서 이미 구멍이 난다.",
        "w_splat_area의 기준이 장면 p99라, 거대 splat도 벌점 0.  (로그에 매 스텝 0)",
        "coverage 지분 ≈ 0.02%.  매칭 항(centroid+chamfer+OT)이 손실의 63%.",
        "전역 scale 밴드는 칸 크기와 무관 — 작은 칸도 큰 splat이 된다.",
    ], 14, INK, 5)
    rect(s, Inches(6.8), Inches(3.3), Inches(6.1), Inches(3.45), PALE)
    txt(s, Inches(7.0), Inches(3.45), Inches(5.75), Inches(0.38), "이 구조에서 나오는 그림", 16, True, NAVY)
    bullets(s, Inches(7.0), Inches(3.95), Inches(5.7), Inches(2.6), [
        "중복된 점은 Chamfer가 칭찬한다 (nn_unique ≈ 0.51).",
        "빈 픽셀은 옆 splat이 커지며 채운다.",
        "가로+세로 획이 원형 가우시안으로 평균된다.",
        "사진 PSNR은 올라가더라도 글자는 안 산다.",
    ], 14, INK, 5)

    # ---- 6 개선 방향 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "2. 개선", "latent·K·디코더 개수는 유지.  구멍과 splat shortcut만 막는다 (fix_6).")
    foot(s, 6)
    bullets(s, Inches(0.5), Inches(1.22), Inches(12.3), Inches(2.15), [
        "① 개수: 인코더 칸당 512점 (overflow 0.0–0.2%).  출력 262k는 그대로.  구멍은 coverage / p2g / radius로 점을 보내 메운다.",
        "② 외형: 위치 대응만 쓰지 않고 칸 안 속성 집합 거리(w_attr_set).  rot 축 이름 손실은 끄고, 화면 타원(w_cov3d)만 맞춘다.",
        "③ 렌더: xyz는 계속 detach.  scale은 칸 중간 크기 기준 밴드로 묶고, 렌더는 1200부터 낮은 해상도로 켠다.",
    ], 15, INK, 6)
    captioned_pic(s, IMG / "r8h_028630_own.png", Inches(0.7), Inches(3.6),
                  Inches(5.7), Inches(3.15), "R8H  ·  같은 계약, 구멍→큰 splat")
    captioned_pic(s, IMG / "p1_028630_own.png", Inches(6.9), Inches(3.6),
                  Inches(5.7), Inches(3.15), "이후 런  ·  실루엣은 남음, 세부는 아직")

    # ---- 7 개선 항목 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "2. 개선  —  코드에서 바뀐 것", "디코더 262,144 / z 32×64×64 는 그대로.  한도와 손실만.")
    foot(s, 7)
    rows = [
        ("어디", "R8H (같은 구조)", "fix_6 / unified"),
        ("인코더 입력", "칸당 144  →  overflow 14–23%", "칸당 512  →  overflow ≈ 0"),
        ("scale", "전역 밴드. 작은 칸도 큰 splat", "칸 중간 크기 기준 ±3"),
        ("속성 읽기", "appearance만. shape는 렌더에 0", "shape를 zero-init 분기로 더함"),
        ("매칭 63%", "centroid+chamfer+OT가 손실 대부분", "0.3 / 0.5 / 0.8 로 축소"),
        ("구멍", "coverage 지분 0.02%", "w_coverage=1200, p2g=3, radius=4"),
        ("속성 다양성", "위치 키 → 방향 다양성 2.3%", "w_attr_set=3, w_rot=0, w_cov3d=2"),
        ("splat 벌점", "p99 기준이라 로그 splat=0", "끔. 분포가 이미 0.40× 너무 작음"),
        ("렌더 커리큘럼", "늦게 켬", "1200부터, 해상도 ×8→×2"),
    ]
    y = Inches(1.18)
    ws = [Inches(2.15), Inches(5.2), Inches(5.2)]
    for r, row in enumerate(rows):
        x = Inches(0.4)
        hh = Inches(0.55)
        for c, (cell, ww) in enumerate(zip(row, ws)):
            if r == 0:
                bg, fg, bld = NAVY, WHITE, True
            elif c == 1:
                bg, fg, bld = RED_BG, INK, False
            elif c == 2:
                bg, fg, bld = GREEN_BG, INK, False
            else:
                bg, fg, bld = PALE, NAVY, True
            rect(s, x, y, ww - Inches(0.05), hh, bg, LINE if r else None)
            txt(s, x + Inches(0.08), y, ww - Inches(0.16), hh,
                cell, 11 if r else 12, bld, fg, "left", "middle")
            x += ww
        y += hh + Inches(0.04)

    # ---- 8 예정 + 정리 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "3. 앞으로 진행할 예정", "계약은 유지한 채, 구조가 아직 남긴 구멍.")
    foot(s, 8)
    triples = [
        (RED, "아직 구조에 남은 것",
         "출력은 계속 262k — N이 200·300만이 되면 부분집합 천장이 내려간다.\n"
         "appearance 8채널 / 64슬롯은 latent 계약이라 그대로다.\n"
         "나중 런에서도 nn_unique ≈ 0.51.  중복은 손실이 아직 안 잡는다.\n"
         "사진 PSNR이 GT보다 높아지는 경우 = 과다평활이 PSNR을 이긴다."),
        (GREEN, "그래서 할 일",
         "262k 타깃을 ‘임의 부분집합’이 아니라\n렌더가 같은 축소 표현으로 만든다.\n"
         "초반에 scale shortcut이 자리 잡기 전에\n색·방향 다양성을 먼저 연다.\n"
         "coverage + attr_set 을 from-scratch 한 런에서 같이 검증한다."),
        (GOLD, "데이터 / 판정",
         "train+truck, vanilla dump.  씬별 앵커\n(공유 앵커는 train 칸 43.7% 공실).\n"
         "판정: codec_own 글자 가독성,\nPSNR gap, SSIM, nn_unique.\n"
         "씬이 커질수록 262k vs 사진 PSNR이\n떨어지면 타깃을 렌더-보존 축소로 교체."),
    ]
    for i, (c, t, b) in enumerate(triples):
        x = Inches(0.4 + i * 4.25)
        rect(s, x, Inches(1.28), Inches(4.05), Inches(5.5), PALE)
        rect(s, x, Inches(1.28), Inches(4.05), Inches(0.7), c)
        txt(s, x, Inches(1.28), Inches(4.05), Inches(0.7), t, 18, True, WHITE, "center", "middle")
        txt(s, x + Inches(0.2), Inches(2.15), Inches(3.65), Inches(4.4), b, 14, False, INK)

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()

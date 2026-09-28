#!/usr/bin/env python3
"""R8H blur problem vs fix_6 — Korean 16:9 deck."""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn, nsmap
from pptx.util import Emu, Inches, Pt
from lxml import etree
from copy import deepcopy

OUT = "/data/daeho/aacd_proj/can3tok_encoder_decoder_new_fix_6/R8H_blur_to_fix6.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
NAVY2 = RGBColor(0x24, 0x3B, 0x6B)
RED = RGBColor(0xC4, 0x3A, 0x31)
RED_BG = RGBColor(0xFB, 0xE9, 0xE6)
GREEN = RGBColor(0x1F, 0x6B, 0x4A)
GREEN_BG = RGBColor(0xE3, 0xF2, 0xEA)
AMBER = RGBColor(0xB4, 0x6C, 0x1A)
AMBER_BG = RGBColor(0xFB, 0xF0, 0xD9)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x64, 0x72)
LINE = RGBColor(0xD7, 0xD2, 0xC8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF4, 0xEE)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SKY = RGBColor(0x2B, 0x5F, 0x8A)

W, H = Inches(13.333), Inches(7.5)
FONT = "Malgun Gothic"


def _set_ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _line(shape, color, pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def rect(slide, x, y, w, h, fill, line=None, radius=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 x, y, w, h)
    _fill(shp, fill)
    if line:
        _line(shp, line, 1.0)
    else:
        shp.line.fill.background()
    if radius is not None:
        shp.adjustments[0] = radius
    return shp


def txt(slide, x, y, w, h, text, size=18, bold=False, color=INK, align="left", anchor="top"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                 "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                   "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    _set_ea(run)
    return box


def add_runs(paragraph, parts, size=16, color=INK, align="left"):
    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
    if paragraph.runs:
        paragraph.runs[0].text = ""
    for t, *opt in parts:
        run = paragraph.add_run()
        run.text = t
        run.font.size = Pt(opt[0] if opt else size)
        run.font.bold = opt[1] if len(opt) > 1 else False
        run.font.color.rgb = opt[2] if len(opt) > 2 else color
        run.font.name = FONT
        _set_ea(run)


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = FONT
        _set_ea(run)
    return box


def header(slide, title, subtitle=None):
    rect(slide, 0, 0, W, H, PAPER)
    rect(slide, 0, 0, W, Inches(0.92), NAVY)
    rect(slide, 0, Inches(0.92), Inches(0.12), Inches(6.58), RED)
    txt(slide, Inches(0.45), Inches(0.18), Inches(12.4), Inches(0.42),
        title, 24, True, WHITE, "left", "middle")
    if subtitle:
        txt(slide, Inches(0.45), Inches(0.52), Inches(12.4), Inches(0.32),
            subtitle, 13, False, RGBColor(0xC5, 0xD0, 0xE0), "left", "middle")
    txt(slide, Inches(11.6), Inches(7.18), Inches(1.5), Inches(0.22),
        "can3tok  ·  fix_6", 10, False, MUTED, "right")


def footer_num(slide, n, total):
    txt(slide, Inches(0.4), Inches(7.18), Inches(2), Inches(0.22),
        f"{n}  /  {total}", 10, False, MUTED)


def card(slide, x, y, w, h, title, body, accent, bg=CARD):
    rect(slide, x, y, w, h, bg, LINE, 0.08)
    rect(slide, x, y, Inches(0.10), h, accent)
    txt(slide, x + Inches(0.28), y + Inches(0.14), w - Inches(0.4), Inches(0.36),
        title, 15, True, accent)
    box = slide.shapes.add_textbox(x + Inches(0.28), y + Inches(0.50),
                                   w - Inches(0.42), h - Inches(0.62))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    run.font.size = Pt(13)
    run.font.color.rgb = INK
    run.font.name = FONT
    _set_ea(run)


def pill(slide, x, y, w, h, text, fill, fg=WHITE):
    shp = rect(slide, x, y, w, h, fill, radius=0.5)
    txt(slide, x, y, w, h, text, 12, True, fg, "center", "middle")
    return shp


def new(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    N = 12

    # 1 title
    s = new(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, Inches(0.18), H, RED)
    txt(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.4),
        "R8H  →  fix_6", 16, False, RGBColor(0xF0, 0xA8, 0x98))
    txt(s, Inches(0.7), Inches(2.15), Inches(12), Inches(1.3),
        "왜 가우시안이 크고 흐렸는가,\n그리고 그 길을 어떻게 막았는가", 32, True, WHITE)
    txt(s, Inches(0.7), Inches(3.7), Inches(11.5), Inches(0.9),
        "디코더를 새로 짠 것이 아니다.\n빠진 표면을 큰 타원 하나로 메우는 쪽이, 점을 제대로 놓는 쪽보다 렌더 손실이 쌌다.",
        16, False, RGBColor(0xC5, 0xD0, 0xE0))
    pill(s, Inches(0.7), Inches(5.5), Inches(2.4), Inches(0.42), "문제  ·  R8H blur", RED)
    pill(s, Inches(3.3), Inches(5.5), Inches(2.6), Inches(0.42), "개선  ·  fix_6 손실·밴드", GREEN)
    txt(s, Inches(0.7), Inches(6.7), Inches(10), Inches(0.3),
        "can3tok encoder–decoder   ·   출력 262,144  ·   z_compact 32×64×64", 12, False,
        RGBColor(0x8A, 0x9B, 0xB5))

    # 2 one-pager
    s = new(prs)
    header(s, "한 장으로", "발표의 결론을 먼저 둔다")
    footer_num(s, 2, N)
    # three columns
    cols = [
        (RED, RED_BG, "문제",
         "R8H codec_own 그림이 큰 splat으로 번진다.\n기관차 실루엣은 남고, 글자·난간·바퀴 살은 녹는다."),
        (AMBER, AMBER_BG, "원인",
         "위치가 비운 구멍을 렌더가 큰 타원으로 메우는 길이 더 싸다.\n점 손실은 그 구멍을 거의 보지 못했다."),
        (GREEN, GREEN_BG, "개선",
         "한 점이 커질 수 있는 한도를 칸 크기에 묶고,\n빈 표면은 점 손실로 직접 벌하고,\n가로·세로는 원으로 평균 내지 않는다."),
    ]
    for i, (acc, bg, t, b) in enumerate(cols):
        x = Inches(0.4 + i * 4.25)
        rect(s, x, Inches(1.25), Inches(4.05), Inches(5.55), bg, radius=0.06)
        rect(s, x, Inches(1.25), Inches(4.05), Inches(0.62), acc)
        txt(s, x, Inches(1.25), Inches(4.05), Inches(0.62), t, 20, True, WHITE, "center", "middle")
        txt(s, x + Inches(0.25), Inches(2.1), Inches(3.55), Inches(4.4), b, 16, False, INK)

    # 3 evidence
    s = new(prs)
    header(s, "무엇이 문제였나", "비교 PNG 다섯 칸이 원인을 위치와 크기로 가른다")
    footer_num(s, 3, N)
    tiles = [
        ("orig", "GT 전부\nGT 크기", "원본 3DGS"),
        ("canon", "고른 26.2만\nGT 크기", "고르기만 해도 남는가"),
        ("codec", "모델 위치\n근처 GT 크기", "위치만의 잘못"),
        ("snap", "최근접 GT\n모델 크기", "위치만 붙이면"),
        ("codec_own", "모델 위치\n모델 크기", "실제로 배포되는 그림"),
    ]
    for i, (name, mid, q) in enumerate(tiles):
        x = Inches(0.35 + i * 2.58)
        fill = RED if name == "codec_own" else (AMBER_BG if name == "codec" else CARD)
        fg = WHITE if name == "codec_own" else INK
        acc = RED if name == "codec_own" else (AMBER if name == "codec" else NAVY)
        rect(s, x, Inches(1.2), Inches(2.42), Inches(2.55), fill, LINE, 0.08)
        txt(s, x, Inches(1.28), Inches(2.42), Inches(0.4), name, 14, True, acc if name != "codec_own" else WHITE, "center")
        txt(s, x + Inches(0.1), Inches(1.7), Inches(2.22), Inches(0.9), mid, 13, False,
            fg, "center")
        txt(s, x + Inches(0.1), Inches(2.65), Inches(2.22), Inches(0.9), q, 12, False,
            WHITE if name == "codec_own" else MUTED, "center")

    # metrics table as cards
    txt(s, Inches(0.4), Inches(3.95), Inches(12), Inches(0.35),
        "R8H  ·  step 10,000  ·  원본 가우시안 렌더 대비", 14, True, NAVY)
    rows = [
        ("장면", "codec  위치만", "codec_own  모델 크기", "서로 다른 GT에 붙은 점"),
        ("018400  기관차", "17.16 dB   SSIM 0.57", "16.24 dB   SSIM 0.43", "48%"),
        ("028630", "18.80 dB   SSIM 0.67", "19.03 dB   SSIM 0.59", "57%"),
    ]
    yw = Inches(4.35)
    widths = [Inches(2.6), Inches(3.3), Inches(3.5), Inches(2.8)]
    for r, row in enumerate(rows):
        x0 = Inches(0.4)
        for c, (cell, ww) in enumerate(zip(row, widths)):
            bg = NAVY if r == 0 else (RED_BG if c == 2 and r > 0 else CARD)
            fg = WHITE if r == 0 else (RED if c == 2 else INK)
            rect(s, x0, yw, ww - Inches(0.06), Inches(0.48), bg, LINE if r else None)
            txt(s, x0, yw, ww - Inches(0.06), Inches(0.48), cell, 13, r == 0 or c == 0, fg, "center", "middle")
            x0 += ww
        yw += Inches(0.52)
    txt(s, Inches(0.4), Inches(6.55), Inches(12.4), Inches(0.5),
        "codec 칸에는 차체 구멍이 보인다. 같은 위치에 모델 scale을 씌운 codec_own은 그 구멍이 큰 페인트가 된다.\n"
        "028630은 own PSNR이 조금 높아도 SSIM은 더 낮다. 흐린 그림이 점수상 괜찮아 보이는 이유다.",
        13, False, MUTED)

    # 4 mechanism
    s = new(prs)
    header(s, "핵심 원인", "빠진 표면을 큰 타원 하나로 메우는 쪽이 더 싸다")
    footer_num(s, 4, N)
    txt(s, Inches(0.45), Inches(1.2), Inches(12.4), Inches(0.7),
        "차체 한 조각에 GT 가우시안이 80개 있다고 하자.\n모델 점은 40개만 그 근처에 있고, 나머지는 이미 덮인 모서리에 중복된다.  (nn_unique ≈ 0.51)",
        16, False, INK)

    # two paths
    rect(s, Inches(0.4), Inches(2.15), Inches(6.05), Inches(4.55), RED_BG, radius=0.06)
    txt(s, Inches(0.6), Inches(2.3), Inches(5.7), Inches(0.4), "길 B  —  렌더가 고른 길", 18, True, RED)
    bullets(s, Inches(0.65), Inches(2.8), Inches(5.6), Inches(3.6), [
        "옆 점의 log_scale을 조금만 올린다.",
        "exp 때문에 화면 타원이 금방 커진다.",
        "opacity를 올리면 타원이 진해진다.",
        "한 점이 80개 몫의 픽셀을 덮는다.",
        "렌더 손실이 당장 줄어든다. 최적화가 여기를 고른다.",
    ], 15, INK, 10)

    rect(s, Inches(6.85), Inches(2.15), Inches(6.05), Inches(4.55), CARD, LINE, 0.06)
    txt(s, Inches(7.05), Inches(2.3), Inches(5.7), Inches(0.4), "길 A  —  점을 옮긴다  (비싸다)", 18, True, NAVY)
    bullets(s, Inches(7.1), Inches(2.8), Inches(5.6), Inches(3.6), [
        "나머지 40개를 빈 곳으로 밀어야 한다.",
        "Chamfer는 모서리에 붙은 점을 이미 ‘맞다’고 친다.",
        "xyz는 detach — 화면 그래디언트가 위치로 거의 안 간다.",
        "발자국 밖 픽셀의 위치 그래디언트는 엉뚱한 방향이다.",
        "측정: 위치 vs 진짜 보정 코사인 0.13, 점의 24%만 그래디언트.",
    ], 15, INK, 10)

    # 5 why matching failed
    s = new(prs)
    header(s, "점 손실이 구멍을 못 본 이유", "대칭 Chamfer는 중복을 칭찬한다")
    footer_num(s, 5, N)
    card(s, Inches(0.4), Inches(1.25), Inches(6.2), Inches(2.6),
         "대칭 Chamfer가 하는 일",
         "예측→GT, GT→예측 최근접의 평균이다.\n예측 점이 GT 모서리에만 잔뜩 붙어 있으면 예측→GT는 짧다. GT 차체 중앙의 빈 쪽은 잘 안 보인다.\n두 예측이 같은 GT에 붙어도 ‘맞다’.",
         AMBER)
    card(s, Inches(6.8), Inches(1.25), Inches(6.1), Inches(2.6),
         "실측",
         "R8H eval  nn_unique ≈ 0.51\n예측 점 둘 중 하나는 이미 다른 점이 맡은 GT에 또 붙는다.\n표면의 반은 비어 있다.\n018400 장면에서는 48%.",
         RED, RED_BG)
    card(s, Inches(0.4), Inches(4.05), Inches(6.2), Inches(2.55),
         "손실의 63%가 ‘고른 26.2만 따라가기’",
         "그룹 중심 4.0 + 그룹 안 Chamfer 6.0 + Sinkhorn 8.0.\n고른 부분집합의 통계에 세게 맞추면, 안 고른 표면을 그림에서 덮으라는 신호는 거의 없다.",
         NAVY)
    card(s, Inches(6.8), Inches(4.05), Inches(6.1), Inches(2.55),
         "coverage는 사실상 꺼져 있었다",
         "‘GT에서 예측까지’ 한쪽 Chamfer가 구멍을 보는 항이다.\nR8H에서 이 항의 손실 지분은 약 0.02%.\n빈 표면을 벌할 목소리가 없었다.",
         MUTED)

    # 6 dead anti-splat + orientation
    s = new(prs)
    header(s, "막으려던 항이 안 걸렸고, 획은 원으로 평균됐다", "두 가지가 같이 흐림을 만든다")
    footer_num(s, 6, N)
    rect(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(5.55), RED_BG, radius=0.06)
    txt(s, Inches(0.6), Inches(1.35), Inches(5.85), Inches(0.4), "큰 splat 금지 항은 죽어 있었다", 17, True, RED)
    bullets(s, Inches(0.65), Inches(1.9), Inches(5.7), Inches(4.5), [
        "w_splat_area=2 가 켜져 있었다.",
        "기준이 장면 전체 99백분위.",
        "차체를 덮는 거대 splat도 배경 큰 splat 꼬리 안에 들어간다.",
        "로그 splat=0.00000 이 매 스텝 — 벌점이 한 번도 안 걸림.",
        "전역 tanh ±3 이면 이론상 약 20배까지 커질 수 있다.",
        "구멍 메우기에는 충분하고, p99에는 안 걸린다.",
    ], 15, INK, 8)

    rect(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.55), AMBER_BG, radius=0.06)
    txt(s, Inches(7.0), Inches(1.35), Inches(5.75), Inches(0.4), "가로 + 세로 = 둥근 평균", 17, True, AMBER)
    bullets(s, Inches(7.05), Inches(1.9), Inches(5.65), Inches(4.5), [
        "한 자리에 세로로 얇은 GT와 가로로 얇은 GT.",
        "슬롯 배정·Sinkhorn·responsibility는 전부 위치만 본다.",
        "크기는 짝 비용에 안 들어간다.",
        "타깃은 둘의 평균 → 둥근 가우시안.",
        "모델이 둥글게 그린 것은 버그가 아니라 질문에 대한 정답.",
        "방향 다양성 = GT의 2.3%.  회전 nrmse 1.12 ≈ 평균만 냄.",
        "글자·난간이 녹는 이유.",
    ], 15, INK, 6)

    # 7 third cause
    s = new(prs)
    header(s, "처음부터 점이 부족했다", "인코더가 이미 14–23%를 버리고 들어온다")
    footer_num(s, 7, N)
    items = [
        ("칸당 144점 상한", "R8H max_input = 589,824.\n밀집 스냅샷에서 칸 overflow 14–23%.\n디코더는 없는 기하를 복원할 수 없다. 렌더만 남고 다시 큰 splat."),
        ("출력은 항상 26.2만", "원본이 더 많아도 디코더는 64점×4096칸.\n고른 부분집합이 원본을 못 대표하면, 빈 곳은 그림에서만 메워진다."),
        ("렌더는 늦게 켜짐", "R8H render_start = 2500.\n점 손실만으로 오래 가면 중복 최적에 먼저 들어간다.\n그 다음 렌더가 구멍을 큰 타원으로 메운다."),
    ]
    for i, (t, b) in enumerate(items):
        y = Inches(1.25 + i * 1.85)
        rect(s, Inches(0.4), y, Inches(12.5), Inches(1.7), CARD, LINE, 0.06)
        rect(s, Inches(0.4), y, Inches(0.12), Inches(1.7), AMBER)
        txt(s, Inches(0.8), y + Inches(0.2), Inches(11.8), Inches(0.4), t, 18, True, NAVY)
        txt(s, Inches(0.8), y + Inches(0.65), Inches(11.8), Inches(0.9), b, 15, False, INK)

    # 8 strategy
    s = new(prs)
    header(s, "fix_6 전략", "디코더는 그대로 두고, 그 구멍에서 싼 길을 바꾼다")
    footer_num(s, 8, N)
    txt(s, Inches(0.45), Inches(1.15), Inches(12.4), Inches(0.55),
        "층수, 잠재 32×64×64, 출력 262,144는 R8H와 같다. 바꾼 것은 한 점이 커질 수 있는 한도와, ‘맞다’고 부르는 손실이다.",
        15, False, MUTED)
    rows = [
        (GREEN, "1. 못 키우게", "칸의 중간 크기를 중심으로 scale을 묶는다. 차체 칸의 점이 배경만큼 커질 수 없다."),
        (SKY, "2. 점으로 메우게", "빈 GT를 coverage / p2g / radius가 직접 벌한다. 큰 splat은 이 항에 진다."),
        (NAVY2, "3. 획을 살리는 타깃", "속성 집합 거리. 가로·세로를 원으로 합치지 않는다. 화면 타원(공분산)만 맞춘다."),
        (AMBER, "4. 구멍을 덜 만든 채 렌더", "칸당 512점. 렌더는 1200부터, 흐린 해상도로. 위치는 계속 detach."),
    ]
    for i, (c, t, b) in enumerate(rows):
        y = Inches(1.85 + i * 1.2)
        rect(s, Inches(0.4), y, Inches(12.5), Inches(1.08), CARD, LINE, 0.08)
        pill(s, Inches(0.55), y + Inches(0.3), Inches(2.35), Inches(0.48), t, c)
        txt(s, Inches(3.15), y + Inches(0.22), Inches(9.5), Inches(0.7), b, 16, False, INK, "left", "middle")

    # 9 scale band
    s = new(prs)
    header(s, "개선 1  ·  한 점을 20배로 키우기 어렵게", "attr_decoder._bounded_log_scale")
    footer_num(s, 9, N)
    rect(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(5.55), RED_BG, radius=0.06)
    txt(s, Inches(0.6), Inches(1.4), Inches(5.85), Inches(0.4), "R8H", 16, True, RED)
    bullets(s, Inches(0.65), Inches(1.95), Inches(5.7), Inches(4.5), [
        "장면 전체 하나의 중심  (bias ≈ −7.58)",
        "tanh cap ±3  →  약 20배",
        "GT log-scale의 5.94%가 이 밴드 밖",
        "모델 퍼짐은 GT의 0.74배  (다양성 죽음)",
        "작은 건 더 작고, 구멍은 전역 꼬리 안에서 커질 수 있음",
        "appearance 8채널만 봄  →  세 축이 비슷해져 구가 됨",
    ], 15, INK, 8)

    rect(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.55), GREEN_BG, radius=0.06)
    txt(s, Inches(7.0), Inches(1.4), Inches(5.75), Inches(0.4), "fix_6", 16, True, GREEN)
    bullets(s, Inches(7.05), Inches(1.95), Inches(5.65), Inches(4.5), [
        "중심 = 그 칸의 중간 크기",
        "0.787 · log(칸 반지름) − 2.559",
        "그 주변 ±3 이면 데이터 99.93%",
        "칸 반지름이 작으면 허용 splat도 작다",
        "초기값도 칸 중간 크기  →  처음부터 거대 타원이 아님",
        "shape 19채널을 zero-init으로 더해 읽음  →  판은 한 축만 납작",
    ], 15, INK, 8)

    # 10 coverage
    s = new(prs)
    header(s, "개선 2  ·  빈 표면을 직접 벌한다", "큰 splat은 중심 하나만 가깝고, 나머지 GT는 여전히 멀다")
    footer_num(s, 10, N)
    table = [
        ("항", "R8H", "fix_6", "구멍에서 하는 일"),
        ("coverage", "지분 0.02%", "1200  ·  4.9만 샘플", "안 덮인 GT가 직접 비용"),
        ("p2g", "없음", "3.0", "허공에 뜬 타원을 벌함"),
        ("radius", "없음", "4.0", "오그라든 칸(0.67×→0.91×)을 펼침"),
        ("점 매칭", "cen/ich/ot = 63%", "약 1/10", "중복을 ‘맞다’고 하던 목소리를 줄임"),
    ]
    y = Inches(1.25)
    ws = [Inches(2.2), Inches(2.5), Inches(3.1), Inches(4.5)]
    for r, row in enumerate(table):
        x = Inches(0.4)
        for c, (cell, ww) in enumerate(zip(row, ws)):
            if r == 0:
                bg, fg, bld = NAVY, WHITE, True
            elif c == 1:
                bg, fg, bld = RED_BG, RED, False
            elif c == 2:
                bg, fg, bld = GREEN_BG, GREEN, True
            else:
                bg, fg, bld = CARD, INK, False
            rect(s, x, y, ww - Inches(0.06), Inches(0.72), bg, LINE if r else None)
            txt(s, x + Inches(0.08), y, ww - Inches(0.14), Inches(0.72), cell, 13, bld, fg, "center", "middle")
            x += ww
        y += Inches(0.78)
    txt(s, Inches(0.45), Inches(5.4), Inches(12.4), Inches(1.5),
        "커버리지는 이런 측정에서 나왔다. GT의 절반만 그려도 19.7 dB, 전체면 48 dB인데\n"
        "모델은 서로 다른 점이 약 48%였다. 절반을 큰 splat으로 메우면 사진 PSNR은 나와도, 그게 R8H blur다.\n"
        "점 40개를 빈 곳에 뿌리는 쪽이 coverage에는 이기고, 한 점을 키우는 쪽은 진다.",
        15, False, MUTED)

    # 11 attr set + curriculum
    s = new(prs)
    header(s, "개선 3–4  ·  획을 살리고, 구멍을 덜 만든 채 렌더를 켠다", "")
    footer_num(s, 11, N)
    card(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(3.35),
         "속성 집합  ·  w_attr_set = 3",
         "칸 안 속성 집합만 비교한다. 위치 짝을 안 맺는다.\n세로 얇은 것과 가로 얇은 것이 있으면 타깃은 ‘둘 다 있는 집합’이다. 둥근 평균은 집합 거리가 크다.\n방향 다양성 2.3% → 15.8%.  PSNR 17.80 → 18.67.\n회전 직접 손실은 끄고, 화면 타원 R diag(s²) Rᵀ 만 맞춘다 (w_cov3d=2).",
         GREEN, GREEN_BG)
    card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(3.35),
         "w_splat_area 는 끈다",
         "R8H에서 이미 매 스텝 0이었다.\n살아 있어도 평균 splat이 GT의 0.40배로 너무 작은 쪽을 더 밀었다.\np80 힌지는 코드에 남아 있지만 unified 레시피는 0.",
         MUTED)
    card(s, Inches(0.4), Inches(4.7), Inches(6.2), Inches(1.95),
         "인코더 칸당 512",
         "overflow 14–23% → 0.0–0.2%.\n입력에서 구멍을 덜 만든다. pool_chunk를 자동으로 둬 메모리는 거의 그대로.",
         SKY)
    card(s, Inches(6.8), Inches(4.7), Inches(6.1), Inches(1.95),
         "렌더 1200, 위치는 detach",
         "너무 늦으면 9.5 dB 구덩이. 너무 이르고 위치까지 밀면 큰 splat이 이긴다.\n겉모습은 일찍, 위치는 발자국 안 nudge만. 파라미터 항 floor 0.5.",
         NAVY)

    # 12 before after + what stayed
    s = new(prs)
    header(s, "같은 구멍, 그리고 정리", "디코더는 같고, 그 구멍에서 최적점이 바뀌었다")
    footer_num(s, 12, N)
    rect(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(3.35), RED_BG, radius=0.06)
    txt(s, Inches(0.6), Inches(1.35), Inches(5.85), Inches(0.35), "R8H", 16, True, RED)
    bullets(s, Inches(0.6), Inches(1.75), Inches(5.8), Inches(2.6), [
        "점 40개가 모서리에 중복  (Chamfer: 맞음)",
        "차체 중앙은 빈 픽셀",
        "렌더가 옆 점 scale을 키움",
        "splat 벌점은 p99라 0",
        "가로·세로 획은 둥근 평균",
        "codec_own = 흐린 페인트",
    ], 14, INK, 4)

    rect(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(3.35), GREEN_BG, radius=0.06)
    txt(s, Inches(7.0), Inches(1.35), Inches(5.75), Inches(0.35), "fix_6", 16, True, GREEN)
    bullets(s, Inches(7.0), Inches(1.75), Inches(5.7), Inches(2.6), [
        "입력에서 점을 덜 버림",
        "coverage가 빈 GT를 직접 벌함",
        "scale이 칸 중간 크기에 묶임",
        "attr_set이 획을 원으로 안 합침",
        "렌더는 작은 타원의 색·투명도만",
        "구멍은 점으로, 획은 얇은 타원으로",
    ], 14, INK, 4)

    txt(s, Inches(0.45), Inches(4.7), Inches(12.4), Inches(0.35),
        "그대로 둔 것", 14, True, NAVY)
    txt(s, Inches(0.45), Inches(5.05), Inches(12.4), Inches(0.7),
        "folding 디코더  ·  template 슬롯  ·  출력 262,144  ·  z_compact 32×64×64  ·  gen 없음  ·  xyz detach",
        14, False, INK)
    txt(s, Inches(0.45), Inches(5.75), Inches(12.4), Inches(0.35),
        "바꾼 것", 14, True, GREEN)
    txt(s, Inches(0.45), Inches(6.1), Inches(12.4), Inches(0.85),
        "인코더가 보는 밀도 (144→512)  ·  scale이 커질 수 있는 범위 (전역 밴드→칸 상대 밴드)\n"
        "맞다고 부르는 손실 (부분집합 63% → 커버리지·집합·구조)  ·  렌더를 켜는 시점 (2500→1200)",
        14, False, INK)

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
